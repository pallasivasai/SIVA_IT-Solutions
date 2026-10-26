from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


def build_leads_workbook(rows: list[dict[str, Any]], project: str, run_id: str) -> bytes:
    """Create a real Excel workbook from sourced public search results, never inventing contacts."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sourced Leads"
    sheet.append(["Company / Result", "Public Source URL", "Evidence / Search Snippet", "Contact Status", "Product", "Research Run ID", "Researched At (UTC)"])
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    seen: set[str] = set()
    for row in rows:
        url = str(row.get("Public Source URL", "")).strip()
        title = str(row.get("Company / Result", "")).strip()
        if not title or not url.startswith(("https://", "http://")) or url in seen:
            continue
        seen.add(url)
        sheet.append([title[:300], url[:1000], str(row.get("Evidence / Search Snippet", ""))[:1200],
                      str(row.get("Contact Status", "Needs verification")), project[:200], run_id[:100], now])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="12314A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate([34, 58, 70, 46, 28, 28, 24], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        url_cell = row[1]
        if str(url_cell.value).startswith(("https://", "http://")):
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"
    if sheet.max_row >= 2:
        table = Table(displayName="SourcedLeads", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    notes = workbook.create_sheet("Read Me")
    notes.append(["SIVA IT Solutions — Sourced Lead Research"])
    notes.append(["Product", project])
    notes.append(["Run ID", run_id])
    notes.append(["Created (UTC)", now])
    notes.append(["Important", "These are public search results, not verified sales-qualified leads. Verify company fit and contact details before outreach."])
    notes.append(["Privacy", "No personal email addresses are guessed; no outreach was sent."])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 110
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    notes["A1"].font = Font(bold=True, size=14, color="12314A")
    notes.row_dimensions[5].height = 42
    notes.row_dimensions[6].height = 32
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def save_leads_workbook(content: bytes, project: str, run_id: str, root: str = "data/leads") -> str:
    """Save a local artifact copy; repository copy is published through a reviewable PR."""
    slug = re.sub(r"[^A-Za-z0-9-]+", "-", project.lower()).strip("-")[:32] or "product"
    safe_run = re.sub(r"[^A-Za-z0-9-]+", "", run_id)[:32] or datetime.now(timezone.utc).strftime("%Y%m%d")
    destination = Path(root) / f"{safe_run}-{slug}-leads.xlsx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(content)
    return str(destination)
